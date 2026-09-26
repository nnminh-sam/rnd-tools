"""DocModel -> .docx (python-docx). Citations become superscript [n] + a Sources list."""

from __future__ import annotations

from pathlib import Path

from .model import Block, DocModel, Run, chart_data, describe_sources, runs_text


def _add_runs(paragraph, runs: list[Run]) -> None:
    from docx.shared import RGBColor

    for r in runs:
        if r.cites:
            run = paragraph.add_run("[" + ",".join(map(str, r.cites)) + "]")
            run.font.superscript = True
            run.font.color.rgb = RGBColor(0x44, 0x55, 0x88)
            continue
        run = paragraph.add_run(r.text)
        run.bold = r.bold or None
        run.italic = r.italic or None
        if r.code:
            run.font.name = "Consolas"


def _style(doc, name: str, fallback: str = "Normal") -> str:
    try:
        doc.styles[name]
        return name
    except KeyError:
        return fallback


def _table(doc, header: list[list[Run]], rows: list[list[list[Run]]]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    width = max([len(header)] + [len(r) for r in rows]) if (header or rows) else 0
    if not width:
        return
    table = doc.add_table(rows=0, cols=width)
    table.style = _style(doc, "Light Grid Accent 1", "Table Grid")
    for r_idx, row in enumerate(([header] if header else []) + rows):
        cells = table.add_row().cells
        for c_idx in range(width):
            runs = row[c_idx] if c_idx < len(row) else []
            p = cells[c_idx].paragraphs[0]
            _add_runs(p, runs)
            text = runs_text(runs).strip()
            if r_idx == 0 and header:
                for run in p.runs:
                    run.bold = True
            elif text and text.replace(",", "").replace(".", "").replace("%", "").replace("-", "").isdigit():
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def render(ws, model: DocModel, out: Path, template: Path | None = None, footer_note: str = "") -> Path:
    import docx
    from docx.shared import Inches, Pt

    if template:
        doc = docx.Document(str(template))
        body = doc.element.body
        for child in list(body):
            if not child.tag.endswith("sectPr"):
                body.remove(child)
    else:
        doc = docx.Document()
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(11)
    doc.core_properties.title = model.title
    if model.meta.get("author"):
        doc.core_properties.author = model.meta["author"]

    if model.meta.get("title"):
        doc.add_paragraph(model.meta["title"], style=_style(doc, "Title"))
        sub = " · ".join(v for v in (model.meta.get("subtitle"), model.meta.get("author"), model.meta.get("date")) if v)
        if sub:
            doc.add_paragraph(sub, style=_style(doc, "Subtitle"))

    fig_no = 0
    for b in model.blocks:
        if b.kind == "heading":
            _add_runs(doc.add_heading(level=min(b.level, 4)), b.runs)
        elif b.kind == "para":
            _add_runs(doc.add_paragraph(), b.runs)
        elif b.kind == "quote":
            _add_runs(doc.add_paragraph(style=_style(doc, "Quote")), b.runs)
        elif b.kind == "list":
            for depth, runs in b.items:
                base = "List Number" if b.ordered else "List Bullet"
                name = base if depth == 0 else f"{base} {min(depth + 1, 3)}"
                _add_runs(doc.add_paragraph(style=_style(doc, name, _style(doc, base))), runs)
        elif b.kind == "table":
            _table(doc, b.header, b.rows)
            doc.add_paragraph()
        elif b.kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(b.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif b.kind == "image":
            path = ws.abspath(b.text)
            if path.exists():
                doc.add_picture(str(path), width=Inches(6))
                if runs_text(b.runs):
                    doc.add_paragraph(runs_text(b.runs), style=_style(doc, "Caption"))
        elif b.kind == "chart":
            fig_no += 1
            _chart(ws, doc, b, fig_no)
        elif b.kind == "rule":
            doc.add_paragraph()

    if model.sources:
        doc.add_heading("Sources", level=1)
        for s, text in zip(model.sources, describe_sources(ws, model.sources)):
            p = doc.add_paragraph()
            p.add_run(f"[{s.number}] ").bold = True
            p.add_run(text)
            if s.quote:
                p.add_run(f" — “{s.quote}”").italic = True
            p.paragraph_format.space_after = Pt(2)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = footer_note
    for run in footer.runs:
        run.font.size = Pt(8)
    doc.save(str(out))
    return out


def _chart(ws, doc, b: Block, fig_no: int) -> None:
    """Word has no python-docx chart API: embed a PNG when matplotlib is available,
    always followed by the exact data table so no figure is only visual."""
    from docx.shared import Inches

    cats, series, title = chart_data(ws, b.spec)
    try:
        import matplotlib

        matplotlib.use("Agg")
        import tempfile

        import matplotlib.pyplot as plt

        fig, ax = plt.subplots(figsize=(7, 3.6), dpi=150)
        kind = b.spec.get("type", "bar")
        if kind == "pie":
            ax.pie(series[0][1], labels=cats, autopct="%1.0f%%")
        elif kind == "line":
            for name, vals in series:
                ax.plot(cats, vals, marker="o", label=name)
        else:  # "bar" = horizontal bars, "column" = vertical bars (PowerPoint naming)
            n = len(series)
            for k, (name, vals) in enumerate(series):
                pos = [i + (k - (n - 1) / 2) * 0.8 / n for i in range(len(cats))]
                if kind == "bar":
                    ax.barh(pos, vals, height=0.8 / n, label=name)
                else:
                    ax.bar(pos, vals, width=0.8 / n, label=name)
            if kind == "bar":
                ax.set_yticks(range(len(cats)), cats)
            else:
                ax.set_xticks(range(len(cats)), cats)
        if len(series) > 1 and kind != "pie":
            ax.legend()
        ax.set_title(title)
        fig.tight_layout()
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            fig.savefig(tmp.name)
            plt.close(fig)
            doc.add_picture(tmp.name, width=Inches(6))
    except ImportError:
        pass
    header = [[Run(b.spec.get("x") or "category")]] + [[Run(name)] for name, _ in series]
    rows = [[[Run(c)]] + [[Run(_fmt(vals[i]))] for _, vals in series] for i, c in enumerate(cats)]
    _table(doc, header, rows)
    src = f" (source: calc:{b.spec['calc']})" if "calc" in b.spec else ""
    doc.add_paragraph(f"Figure {fig_no}: {title or 'Chart data'}{src}", style=_style(doc, "Caption"))


def _fmt(v: float) -> str:
    from ..util import display_number

    return display_number(v)
