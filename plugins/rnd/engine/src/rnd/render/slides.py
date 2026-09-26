"""DocModel -> .pptx (python-pptx).

Slides are separated by `---`. The first heading of a slide is its title; lists and
paragraphs become bullets; a pipe table becomes a native table; a ```chart block
becomes a native, editable PowerPoint chart built from exact calc data; an HTML
comment `<!-- notes: ... -->` becomes speaker notes. Sources go on closing slides.
"""

from __future__ import annotations

from pathlib import Path

from .model import Block, DocModel, Run, chart_data, describe_sources, runs_text

MAX_BULLET_CHARS = 900


def _split_slides(model: DocModel) -> list[list[Block]]:
    """Explicit `---` breaks win; a Markdown without them gets one slide per # / ## heading."""
    explicit = any(b.kind == "rule" for b in model.blocks)
    slides: list[list[Block]] = [[]]
    for b in model.blocks:
        if b.kind == "rule":
            slides.append([])
            continue
        if not explicit and b.kind == "heading" and b.level <= 2 and slides[-1]:
            slides.append([])
        slides[-1].append(b)
    return [s for s in slides if s]


def _runs_to_paragraph(p, runs: list[Run], size) -> None:
    for r in runs:
        run = p.add_run()
        if r.cites:
            run.text = "[" + ",".join(map(str, r.cites)) + "]"
            run.font.size = size
            rpr = run._r.get_or_add_rPr()
            rpr.set("baseline", "30000")
            continue
        run.text = r.text
        run.font.size = size
        if r.bold:
            run.font.bold = True
        if r.italic:
            run.font.italic = True


def _text_items(blocks: list[Block]) -> list[tuple[int, list[Run]]]:
    items: list[tuple[int, list[Run]]] = []
    for b in blocks:
        if b.kind in ("para", "quote"):
            items.append((0, b.runs))
        elif b.kind == "list":
            items += b.items
        elif b.kind == "heading":
            items.append((0, [Run(runs_text(b.runs), bold=True)]))
        elif b.kind == "code":
            items.append((0, [Run(b.text, code=True)]))
    return items


def _font_size(items):
    from pptx.util import Pt

    chars = sum(len(runs_text(r)) for _, r in items)
    if chars > 700 or len(items) > 9:
        return Pt(13)
    if chars > 450 or len(items) > 7:
        return Pt(15)
    if chars > 250:
        return Pt(18)
    return Pt(20)


def _fill_text(frame, items) -> None:
    size = _font_size(items)
    frame.word_wrap = True
    first = True
    for depth, runs in items:
        p = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        p.level = min(depth, 4)
        _runs_to_paragraph(p, runs, size)


def _add_table(slide, b: Block, left, top, width, height) -> None:
    from pptx.util import Pt

    rows = ([b.header] if b.header else []) + b.rows
    n_cols = max(len(r) for r in rows)
    shape = slide.shapes.add_table(len(rows), n_cols, left, top, width, height)
    size = Pt(14 if len(rows) <= 6 else 12 if len(rows) <= 10 else 10)
    for r_idx, row in enumerate(rows):
        for c_idx in range(n_cols):
            cell = shape.table.cell(r_idx, c_idx)
            cell.text = ""
            runs = row[c_idx] if c_idx < len(row) else []
            _runs_to_paragraph(cell.text_frame.paragraphs[0], runs, size)
            if r_idx == 0 and b.header:
                for run in cell.text_frame.paragraphs[0].runs:
                    run.font.bold = True


def _add_chart(ws, slide, b: Block, left, top, width, height) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

    cats, series, title = chart_data(ws, b.spec)
    data = CategoryChartData()
    data.categories = cats
    for name, vals in series:
        data.add_series(name, vals)
    kind = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
    }.get(b.spec.get("type", "column"), XL_CHART_TYPE.COLUMN_CLUSTERED)
    chart = slide.shapes.add_chart(kind, left, top, width, height, data).chart
    chart.has_legend = len(series) > 1 or kind == XL_CHART_TYPE.PIE
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title
    if kind != XL_CHART_TYPE.PIE:
        plot = chart.plots[0]
        plot.has_data_labels = True
        integers = all(float(v).is_integer() for _, vals in series for v in vals)
        plot.data_labels.number_format = "#,##0" if integers else "#,##0.0#"
        plot.data_labels.number_format_is_linked = False


def _clear_slides(prs) -> None:
    sld_ids = prs.slides._sldIdLst
    for sld in list(sld_ids):
        prs.part.drop_rel(sld.rId)
        sld_ids.remove(sld)


def render(ws, model: DocModel, out: Path, template: Path | None = None, footer_note: str = "") -> Path:
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt

    prs = Presentation(str(template)) if template else Presentation()
    if template:
        _clear_slides(prs)
    else:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layouts = prs.slide_layouts
    title_layout = layouts[0]
    content_layout = layouts[1] if len(layouts) > 1 else layouts[0]
    title_only = layouts[5] if len(layouts) > 5 else content_layout
    W, H = prs.slide_width, prs.slide_height
    margin = Inches(0.5)
    top = Inches(1.45)
    body_h = H - top - Inches(0.6)

    if model.meta.get("title"):
        s = prs.slides.add_slide(title_layout)
        s.shapes.title.text = model.meta["title"]
        sub = " · ".join(v for v in (model.meta.get("subtitle"), model.meta.get("author"), model.meta.get("date")) if v)
        if sub and len(s.placeholders) > 1:
            s.placeholders[1].text = sub

    for blocks in _split_slides(model):
        title_block = next((b for b in blocks if b.kind == "heading"), None)
        rest = [b for b in blocks if b is not title_block]
        visuals = [b for b in rest if b.kind in ("table", "chart", "image")]
        texts = _text_items([b for b in rest if b.kind not in ("table", "chart", "image", "notes")])
        notes = "\n".join(b.text for b in rest if b.kind == "notes")

        slide = prs.slides.add_slide(title_only if visuals else content_layout)
        if slide.shapes.title is not None and title_block is not None:
            slide.shapes.title.text = ""
            _runs_to_paragraph(slide.shapes.title.text_frame.paragraphs[0], title_block.runs, Pt(30))
        body_ph = next((ph for ph in slide.placeholders if ph.placeholder_format.idx == 1), None)

        if visuals:
            if texts:
                text_w = int((W - 2 * margin) * 0.38)
                box = slide.shapes.add_textbox(margin, top, text_w, body_h)
                _fill_text(box.text_frame, texts)
                v_left, v_w = margin + text_w + Inches(0.3), W - 2 * margin - text_w - Inches(0.3)
            else:
                v_left, v_w = margin, W - 2 * margin
            v_h = int(body_h / len(visuals))
            for k, v in enumerate(visuals):
                v_top = top + k * v_h
                if v.kind == "table":
                    _add_table(slide, v, v_left, v_top, v_w, Emu(min(v_h, Inches(0.42) * (len(v.rows) + 1))))
                elif v.kind == "chart":
                    _add_chart(ws, slide, v, v_left, v_top, v_w, v_h)
                elif v.kind == "image":
                    path = ws.abspath(v.text)
                    if path.exists():
                        slide.shapes.add_picture(str(path), v_left, v_top, height=v_h)
        elif texts:
            if body_ph is not None:
                body_ph.text_frame.text = ""
                _fill_text(body_ph.text_frame, texts)
            else:
                box = slide.shapes.add_textbox(margin, top, W - 2 * margin, body_h)
                _fill_text(box.text_frame, texts)
        elif body_ph is not None:
            body_ph._element.getparent().remove(body_ph._element)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        if footer_note:
            fb = slide.shapes.add_textbox(margin, H - Inches(0.45), W - 2 * margin, Inches(0.3))
            fb.text_frame.text = footer_note
            fb.text_frame.paragraphs[0].runs[0].font.size = Pt(9)

    if model.sources:
        entries = describe_sources(ws, model.sources)
        per = 8
        for start in range(0, len(entries), per):
            slide = prs.slides.add_slide(title_only)
            slide.shapes.title.text = "Sources" if start == 0 else "Sources (cont.)"
            box = slide.shapes.add_textbox(margin, top, W - 2 * margin, body_h)
            items = [
                (0, [Run(f"[{s.number}] ", bold=True), Run(text)])
                for s, text in zip(model.sources[start : start + per], entries[start : start + per])
            ]
            _fill_text(box.text_frame, items)
            for p in box.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(12)
    prs.core_properties.title = model.title
    prs.save(str(out))
    return out
