"""Excel (.xlsx/.xlsm) extraction.

Each sheet yields: an info block (size, header, SQL name), preamble rows above the
header, one block per data row ("col: value | ..."), and a Table for SQL.

Values come from Excel's cached results (data_only). A formula cell that has never
been calculated (e.g. written by a script and never opened in Excel) has no cached
value; we count those and warn instead of silently reporting empty cells.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from . import Block, Extraction, Table
from .tabular import col_letter, row_text, sheet_ref

HEADER_SCAN_ROWS = 25
FORMULA_CHECK_MAX_BYTES = 15 * 1024 * 1024


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def detect_header(grid: list[list[Any]]) -> int | None:
    """Index of the header row: the first row whose non-empty cells are all text,
    with at least two of them, covering at least half of the table width."""
    widths = [sum(not _is_blank(v) for v in row) for row in grid[:50]]
    width = max(widths, default=0)
    for i, row in enumerate(grid[:HEADER_SCAN_ROWS]):
        vals = [v for v in row if not _is_blank(v)]
        looks_like_header = len(vals) >= 2 and all(isinstance(v, str) for v in vals) and len(vals) >= 0.5 * width
        if looks_like_header and any(any(not _is_blank(v) for v in r) for r in grid[i + 1 : i + 4]):
            return i
    return None


def _uncalculated_formulas(path: Path) -> dict[str, int]:
    import openpyxl

    counts: dict[str, int] = {}
    if path.stat().st_size > FORMULA_CHECK_MAX_BYTES:
        return counts
    wb_f = openpyxl.load_workbook(path, data_only=False, read_only=True)
    wb_v = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        for ws_f in wb_f.worksheets:
            ws_v = wb_v[ws_f.title]
            n = 0
            for row_f, row_v in zip(ws_f.iter_rows(values_only=True), ws_v.iter_rows(values_only=True)):
                for f, v in zip(row_f, row_v):
                    if v is None and isinstance(f, str) and f.startswith("="):
                        n += 1
            if n:
                counts[ws_f.title] = n
    finally:
        wb_f.close()
        wb_v.close()
    return counts


def extract(path: Path, header_overrides: dict[str, int] | None = None) -> Extraction:
    import openpyxl

    header_overrides = header_overrides or {}
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    blocks: list[Block] = []
    tables: list[Table] = []
    warnings: list[str] = []
    missing = _uncalculated_formulas(path)
    sheets_meta = []
    try:
        title = (wb.properties.title or "").strip() or path.stem
        for ws in wb.worksheets:
            grid: list[list[Any]] = []
            fmts: list[list[str | None]] = []
            for row in ws.iter_rows():
                grid.append([getattr(c, "value", None) for c in row])
                fmts.append([getattr(c, "number_format", None) for c in row])
            # trim empty trailing rows / columns
            while grid and all(_is_blank(v) for v in grid[-1]):
                grid.pop()
                fmts.pop()
            if not grid:
                continue
            ncols = max((max((i + 1 for i, v in enumerate(r) if not _is_blank(v)), default=0) for r in grid), default=0)
            grid = [(r + [None] * ncols)[:ncols] for r in grid]
            fmts = [(r + [None] * ncols)[:ncols] for r in fmts]
            sref = sheet_ref(ws.title)
            last = col_letter(ncols)

            if ws.title in header_overrides:
                h_idx: int | None = header_overrides[ws.title] - 1
            else:
                h_idx = detect_header(grid)

            if h_idx is not None:
                header_cells = grid[h_idx]
                first = next(i for i, v in enumerate(header_cells) if not _is_blank(v))
                header = [
                    str(v).strip() if not _is_blank(v) else f"col_{col_letter(i + 1)}"
                    for i, v in enumerate(header_cells)
                ][first:]
            else:
                first = 0
                header = [f"col_{col_letter(i + 1)}" for i in range(ncols)]

            data_start = (h_idx + 1) if h_idx is not None else 0
            n_data = sum(1 for r in grid[data_start:] if any(not _is_blank(v) for v in r))
            info = (
                f"Sheet {ws.title}: {n_data} data rows x {len(header)} columns"
                + (f"; header on row {h_idx + 1}" if h_idx is not None else "; no header row detected")
                + f". Columns: {', '.join(header)}"
            )
            if ws.title in missing:
                info += f". WARNING: {missing[ws.title]} formula cells have no cached value (open and save in Excel)."
                warnings.append(f"{ws.title}: {missing[ws.title]} uncalculated formula cells")
            blocks.append(Block(f"{sref}!A1:{last}{len(grid)}", "sheet", info, ws.title, {"sheet": ws.title}))

            # rows above the header: titles, notes, units...
            for r in range(h_idx or 0):
                vals = [v for v in grid[r] if not _is_blank(v)]
                if vals:
                    txt = " | ".join(str(v).strip() if isinstance(v, str) else str(v) for v in vals)
                    blocks.append(
                        Block(f"{sref}!A{r + 1}:{last}{r + 1}", "row", txt, ws.title, {"sheet": ws.title, "r": r + 1})
                    )

            body: list[list[Any]] = []
            for r in range(data_start, len(grid)):
                vals = grid[r][first:]
                if all(_is_blank(v) for v in vals):
                    continue
                txt = row_text(header, vals, fmts[r][first:])
                blocks.append(
                    Block(
                        f"{sref}!{col_letter(first + 1)}{r + 1}:{last}{r + 1}",
                        "row",
                        txt,
                        ws.title,
                        {"sheet": ws.title, "r": r + 1, "c0": first + 1},
                    )
                )
                body.append(list(vals))

            if body:
                anchor = f"{sref}!{col_letter(first + 1)}{(h_idx + 1) if h_idx is not None else 1}:{last}{len(grid)}"
                notes = [f"{missing[ws.title]} formula cells without cached values"] if ws.title in missing else []
                tables.append(
                    Table(ws.title, anchor, header, body, (h_idx + 1) if h_idx is not None else None, first + 1, notes)
                )
            sheets_meta.append(
                {
                    "name": ws.title,
                    "rows": len(grid),
                    "cols": ncols,
                    "header_row": (h_idx + 1) if h_idx is not None else None,
                }
            )
    finally:
        wb.close()
    return Extraction("xlsx", title, blocks, tables, warnings, {"sheets": sheets_meta})
