"""PowerPoint (.pptx) extraction: slide text in reading order, tables, chart data, notes."""

from __future__ import annotations

from pathlib import Path

from ..util import format_value
from . import Block, Extraction, Table
from .tabular import coerce_cell, row_text


def _iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _reading_order(shape) -> tuple[int, int]:
    return (shape.top or 0, shape.left or 0)


def extract(path: Path) -> Extraction:
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks: list[Block] = []
    tables: list[Table] = []
    titles: list[str] = []
    for s_no, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        titles.append(title)
        lines: list[str] = [title] if title else []
        t_no = c_no = 0
        for shape in sorted(_iter_shapes(slide.shapes), key=_reading_order):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = para.text.replace("\v", " ").strip()
                    if txt:
                        lines.append(("  " * para.level) + "- " + txt)
            if getattr(shape, "has_table", False) and shape.has_table:
                t_no += 1
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                rows = [r for r in rows if any(r)]
                if rows:
                    header = rows[0]
                    txt = "\n".join([" | ".join(header)] + [row_text(header, r) for r in rows[1:]])
                    blocks.append(
                        Block(f"s{s_no}.t{t_no}", "table_row", txt, title or f"Slide {s_no}", {"slide": s_no})
                    )
                    if len(rows) >= 2 and len(header) >= 2:
                        body = [[coerce_cell(c) for c in r] for r in rows[1:]]
                        tables.append(Table(f"Slide {s_no} table {t_no}", f"s{s_no}.t{t_no}", header, body))
            if getattr(shape, "has_chart", False) and shape.has_chart:
                c_no += 1
                chart = shape.chart
                ctitle = chart.chart_title.text_frame.text.strip() if chart.has_title else ""
                plot = chart.plots[0] if len(chart.plots) else None
                cats = [format_value(c) for c in plot.categories] if plot is not None else []
                series = [(s.name, list(s.values)) for p in chart.plots for s in p.series]
                parts = [f"Chart{': ' + ctitle if ctitle else ''}", "categories: " + ", ".join(cats)]
                parts += [f"{name}: " + ", ".join(format_value(v) for v in vals) for name, vals in series]
                blocks.append(
                    Block(f"s{s_no}.c{c_no}", "chart", " | ".join(parts), title or f"Slide {s_no}", {"slide": s_no})
                )
                if cats and series:
                    cols = ["category"] + [name or f"series{i + 1}" for i, (name, _) in enumerate(series)]
                    body = [
                        [cat] + [vals[i] if i < len(vals) else None for _, vals in series] for i, cat in enumerate(cats)
                    ]
                    tables.append(Table(f"Slide {s_no} chart {c_no}", f"s{s_no}.c{c_no}", cols, body))
        # slide text block goes before its tables/charts
        insert_at = next((i for i, b in enumerate(blocks) if b.meta.get("slide") == s_no), len(blocks))
        blocks.insert(
            insert_at,
            Block(
                f"s{s_no}",
                "slide",
                "\n".join(lines) or f"(slide {s_no} has no text)",
                title or f"Slide {s_no}",
                {"slide": s_no},
            ),
        )
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                blocks.append(Block(f"s{s_no}.notes", "notes", notes, title or f"Slide {s_no}", {"slide": s_no}))

    doc_title = (prs.core_properties.title or "").strip() or next((t for t in titles if t), path.stem)
    return Extraction("pptx", doc_title, blocks, tables, meta={"slides": len(prs.slides)})
