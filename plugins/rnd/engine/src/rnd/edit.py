"""Surgical edits of existing Office files — never rewrite a whole document to change a value.

Every edit writes a new file by default (`<name>.edited.<ext>`); `in_place=True`
keeps a timestamped backup in .rnd/backups/ first. Edits that would silently lose
content (openpyxl drops charts/images when re-saving a workbook) are refused unless
the caller writes to a new file.
"""

from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path
from typing import Any

from .util import now_iso
from .workspace import Workspace


class EditError(RuntimeError):
    pass


def _target(ws: Workspace, src: Path, out: str | None, in_place: bool) -> Path:
    if in_place:
        ws.backups_dir.mkdir(parents=True, exist_ok=True)
        stamp = now_iso().replace(":", "").replace("-", "")
        shutil.copy2(src, ws.backups_dir / f"{src.stem}.{stamp}{src.suffix}")
        return src
    if out:
        return ws.abspath(out)
    return src.with_name(f"{src.stem}.edited{src.suffix}")


def _has_drawings(path: Path) -> bool:
    with zipfile.ZipFile(path) as z:
        return any(n.startswith(("xl/drawings/", "xl/charts/", "xl/media/")) for n in z.namelist())


_CELL = re.compile(r"^(?:(?P<sheet>'(?:[^']|'')+'|[^!]+)!)?(?P<cell>\$?[A-Z]{1,3}\$?\d+)$")


def xlsx_set(ws: Workspace, path: str, cells: dict[str, Any], out: str | None = None, in_place: bool = False) -> str:
    """Set cell values, e.g. {"Quotes!K3": 5.1, "Notes!A1": "checked"}. Formulas start with '='."""
    import openpyxl

    src = ws.abspath(path)
    if not src.exists():
        raise EditError(f"File not found: {path}")
    if in_place and _has_drawings(src):
        raise EditError(
            "Workbook contains charts/images, which openpyxl cannot preserve. Write to a new file (omit in_place)."
        )
    wb = openpyxl.load_workbook(src, keep_vba=src.suffix.lower() == ".xlsm")
    changes = []
    for ref, value in cells.items():
        m = _CELL.match(ref.strip())
        if not m:
            raise EditError(f"Bad cell reference '{ref}' (use Sheet!B4)")
        sheet = (m.group("sheet") or wb.active.title).strip("'").replace("''", "'")
        if sheet not in wb.sheetnames:
            raise EditError(f"Sheet '{sheet}' not in {wb.sheetnames}")
        cell = wb[sheet][m.group("cell").replace("$", "")]
        old = cell.value
        cell.value = value
        changes.append(f"{sheet}!{cell.coordinate}: {old!r} -> {value!r}")
    dst = _target(ws, src, out, in_place)
    wb.save(dst)
    note = (
        " Formula results are computed when the file is opened in Excel."
        if any(isinstance(v, str) and v.startswith("=") for v in cells.values())
        else ""
    )
    return f"Wrote {ws.rel(dst)}:\n  " + "\n  ".join(changes) + note + "\nRe-run index to refresh search and SQL."


def _replace_in_paragraph(paragraph, find: str, repl: str) -> int:
    """Replace text across runs while keeping the formatting of the first run of each match."""
    count = 0
    search_from = 0
    while True:
        runs = paragraph.runs
        full = "".join(r.text for r in runs)
        idx = full.find(find, search_from)
        if idx < 0:
            return count
        count += 1
        search_from = idx + len(repl)  # never re-match inside the replacement
        pos = 0
        start_run = end_run = None
        for i, r in enumerate(runs):
            if start_run is None and pos + len(r.text) > idx:
                start_run, start_off = i, idx - pos
            if start_run is not None and pos + len(r.text) >= idx + len(find):
                end_run, end_off = i, idx + len(find) - pos
                break
            pos += len(r.text)
        if start_run is None or end_run is None:  # pragma: no cover - defensive
            return count
        if start_run == end_run:
            t = runs[start_run].text
            runs[start_run].text = t[:start_off] + repl + t[end_off:]
        else:
            runs[start_run].text = runs[start_run].text[:start_off] + repl
            for k in range(start_run + 1, end_run):
                runs[k].text = ""
            runs[end_run].text = runs[end_run].text[end_off:]


def docx_replace(
    ws: Workspace, path: str, find: str, replace: str, out: str | None = None, in_place: bool = False
) -> str:
    import docx

    if not find:
        raise EditError("find must not be empty")
    src = ws.abspath(path)
    document = docx.Document(str(src))
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs += cell.paragraphs
    for section in document.sections:
        paragraphs += section.header.paragraphs + section.footer.paragraphs
    n = sum(_replace_in_paragraph(p, find, replace) for p in paragraphs)
    if n == 0:
        return f"No occurrences of {find!r} in {path}; nothing written."
    dst = _target(ws, src, out, in_place)
    document.save(str(dst))
    return f"Replaced {n} occurrence(s) of {find!r} in {ws.rel(dst)}. Re-run index to refresh search."


def docx_append_markdown(
    ws: Workspace, path: str, markdown: str, out: str | None = None, in_place: bool = False
) -> str:
    """Append a Markdown section (verified separately) to an existing Word document."""
    import docx

    from .render import model as md_model
    from .render.word import _add_runs, _style, _table

    src = ws.abspath(path)
    document = docx.Document(str(src))
    doc = md_model.parse(markdown)
    for b in doc.blocks:
        if b.kind == "heading":
            _add_runs(document.add_heading(level=min(b.level, 4)), b.runs)
        elif b.kind == "para":
            _add_runs(document.add_paragraph(), b.runs)
        elif b.kind == "list":
            for _depth, runs in b.items:
                base = "List Number" if b.ordered else "List Bullet"
                _add_runs(document.add_paragraph(style=_style(document, base)), runs)
        elif b.kind == "table":
            _table(document, b.header, b.rows)
    if doc.sources:
        entries = md_model.describe_sources(ws, doc.sources)
        document.add_paragraph("Sources for this section:", style=_style(document, "Caption"))
        for s, text in zip(doc.sources, entries):
            document.add_paragraph(f"[{s.number}] {text}")
    dst = _target(ws, src, out, in_place)
    document.save(str(dst))
    return f"Appended {len(doc.blocks)} block(s) to {ws.rel(dst)}."
