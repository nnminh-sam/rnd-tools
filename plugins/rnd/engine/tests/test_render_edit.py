"""Renderers produce valid Office files whose content round-trips through the extractors."""

from rnd.data.tables import run_sql
from rnd.edit import docx_replace, xlsx_set
from rnd.extract import extract
from rnd.render import render_markdown
from rnd.render.sheets import export_calcs

KICKOFF = "2026-02-09-kickoff-project-kestrel"


def _write_md(ws, calc):
    md = f"""---
title: P2 status
---
# Summary

The retail price target is USD 349 [@{KICKOFF}#t1].

---

# Shell fatigue

| Material | Mean cycles |
|---|---|
| rPP-50 | 148,250 [@calc:{calc}] |

```chart
type: column
title: Mean cycles
calc: {calc}
x: material
y: m
```
<!-- notes: speaker note text -->
"""
    path = ws.root / "outputs" / "p2.md"
    path.write_text(md)
    return "outputs/p2.md"


def test_render_all_formats(fresh_ws):
    calc = run_sql(
        fresh_ws,
        "select material, avg(cycles_to_failure) m from "
        "seat_shell_materials_test_results__shell_fatigue group by material order by m desc",
    ).calc_id
    md = _write_md(fresh_ws, calc)
    for fmt in ("docx", "pdf", "pptx"):
        res = render_markdown(fresh_ws, md, fmt)
        assert res.ok, res.message
        assert res.path.exists() and res.path.stat().st_size > 2000
    deck = extract(fresh_ws.root / "outputs" / "p2.pptx")
    anchors = [b.anchor for b in deck.blocks]
    assert "s3.c1" in anchors and "s3.notes" in anchors  # s1 = title slide; native chart + notes survived
    assert any(b.text.startswith("Sources") for b in deck.blocks if b.kind == "slide")
    doc = extract(fresh_ws.root / "outputs" / "p2.docx")
    text = "\n".join(b.text for b in doc.blocks)
    assert "USD 349[1]" in text and "Kickoff Meeting Notes" in text


def test_render_refuses_unverified(fresh_ws):
    (fresh_ws.root / "outputs" / "bad.md").write_text("Sales grew 99% last year.\n")
    res = render_markdown(fresh_ws, "outputs/bad.md", "docx")
    assert not res.ok and "uncited-figure" in res.message
    forced = render_markdown(fresh_ws, "outputs/bad.md", "docx", force=True)
    assert forced.ok and "UNVERIFIED" in forced.message


def test_export_xlsx_has_provenance(fresh_ws):
    import openpyxl

    calc = run_sql(
        fresh_ws, "select model, price_usd from competitor_benchmark_2026__benchmark order by price_usd"
    ).calc_id
    out = fresh_ws.root / "outputs" / "bench.xlsx"
    export_calcs(fresh_ws, [calc], out, ["prices"], chart="bar")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["prices", "Provenance"]
    assert wb["prices"]["B2"].value == 229 and wb["prices"]._charts
    assert "select model" in wb["Provenance"]["F2"].value


def test_xlsx_edit_writes_new_file(fresh_ws):
    msg = xlsx_set(fresh_ws, "sources/competitors/Competitor benchmark 2026.xlsx", {"Benchmark!C2": 239})
    assert "Benchmark!C2: 229 -> 239" in msg
    assert (fresh_ws.root / "sources/competitors/Competitor benchmark 2026.edited.xlsx").exists()


def test_docx_replace_keeps_original(fresh_ws):
    src = "sources/meetings/2026-08-20 Weekly sync.docx"
    msg = docx_replace(fresh_ws, src, "No decision taken.", "Decision: option B.")
    assert msg.startswith("Replaced 1 occurrence")
    edited = extract(fresh_ws.root / "sources/meetings/2026-08-20 Weekly sync.edited.docx")
    assert any("Decision: option B." in b.text for b in edited.blocks)
